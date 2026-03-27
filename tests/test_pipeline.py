from pathlib import Path

import openpyxl

from src.api.pipeline import Pipeline


def _create_test_excel(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Income Statement"

    ws["A1"] = "Metric"
    ws["B1"] = "Q2 2025"
    ws["C1"] = "Q1 2025"

    ws["A2"] = "Revenue"
    ws["B2"] = 1200000
    ws["C2"] = 1100000

    ws["A3"] = "EBITDA"
    ws["B3"] = 300000
    ws["C3"] = 280000

    ws["A4"] = "Net Income"
    ws["B4"] = 150000
    ws["C4"] = 140000

    ws2 = wb.create_sheet("Balance Sheet")
    ws2["A1"] = "Item"
    ws2["B1"] = "Q2 2025"

    ws2["A2"] = "Total Assets"
    ws2["B2"] = 5000000

    ws2["A3"] = "Total Liabilities"
    ws2["B3"] = 3000000

    ws2["A4"] = "Total Equity"
    ws2["B4"] = 2000000

    file_path = path / "test_financials.xlsx"
    wb.save(file_path)
    return file_path


class TestPipeline:
    def test_full_pipeline(self, tmp_path):
        excel_file = _create_test_excel(tmp_path)
        pipeline = Pipeline()
        pipeline.ingestion = __import__(
            "src.ingestion.service", fromlist=["IngestionService"]
        ).IngestionService(upload_dir=str(tmp_path / "uploads"))

        result = pipeline.process_file(
            file_path=excel_file,
            company_id="ACME",
            uploaded_by="test_user",
        )

        assert result.status == "COMPLETED"
        assert result.records_extracted > 0
        assert result.records_created > 0
        assert len(result.actions) > 0

    def test_duplicate_detection(self, tmp_path):
        excel_file = _create_test_excel(tmp_path)
        pipeline = Pipeline()
        pipeline.ingestion = __import__(
            "src.ingestion.service", fromlist=["IngestionService"]
        ).IngestionService(upload_dir=str(tmp_path / "uploads"))

        result1 = pipeline.process_file(excel_file, company_id="ACME")
        result2 = pipeline.process_file(excel_file, company_id="ACME")

        assert result1.status == "COMPLETED"
        assert result2.status == "DUPLICATE"

    def test_pipeline_audit_trail(self, tmp_path):
        excel_file = _create_test_excel(tmp_path)
        pipeline = Pipeline()
        pipeline.ingestion = __import__(
            "src.ingestion.service", fromlist=["IngestionService"]
        ).IngestionService(upload_dir=str(tmp_path / "uploads"))

        result = pipeline.process_file(excel_file, company_id="ACME")

        audit_log = pipeline.audit.get_full_log()
        assert len(audit_log) > 0

        event_types = [e.event_type for e in audit_log]
        assert "FILE_INGESTED" in event_types
        assert "DATA_PARSED" in event_types
