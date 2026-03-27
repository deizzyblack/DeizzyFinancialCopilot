from pathlib import Path

import openpyxl

from src.core.schemas import RawExtraction

FINANCIAL_KEYWORDS = [
    "income",
    "revenue",
    "sales",
    "ebitda",
    "profit",
    "loss",
    "balance",
    "asset",
    "liabilit",
    "equity",
    "cash",
    "expense",
    "operating",
    "financial",
    "statement",
]


def _is_financial_sheet(sheet_name: str) -> bool:
    name_lower = sheet_name.lower()
    return any(kw in name_lower for kw in FINANCIAL_KEYWORDS)


def _cell_ref(row: int, col: int) -> str:
    col_letter = openpyxl.utils.get_column_letter(col)
    return f"{col_letter}{row}"


def _is_numeric(value) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace(" ", "").strip()
        try:
            float(cleaned)
            return True
        except ValueError:
            return False
    return False


def _to_numeric(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace(" ", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


class DocumentParser:
    def parse_excel(self, file_path: Path) -> list[RawExtraction]:
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        extractions = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            if not _is_financial_sheet(sheet_name):
                rows = list(ws.iter_rows(max_row=5, values_only=False))
                has_financial = False
                for row in rows:
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            if any(kw in cell.value.lower() for kw in FINANCIAL_KEYWORDS):
                                has_financial = True
                                break
                    if has_financial:
                        break
                if not has_financial:
                    continue

            headers: dict[int, str] = {}
            rows_list = list(ws.iter_rows(values_only=False))

            if not rows_list:
                continue

            for cell in rows_list[0]:
                if cell.value is not None:
                    headers[cell.column] = str(cell.value)

            for row in rows_list[1:]:
                label_cell = row[0] if row else None
                if label_cell is None or label_cell.value is None:
                    continue

                raw_label = str(label_cell.value).strip()
                if not raw_label:
                    continue

                for cell in row[1:]:
                    if _is_numeric(cell.value):
                        extractions.append(
                            RawExtraction(
                                raw_label=raw_label,
                                raw_value=_to_numeric(cell.value),
                                sheet=sheet_name,
                                cell=_cell_ref(cell.row, cell.column),
                                column_header=headers.get(cell.column),
                                row_index=cell.row,
                            )
                        )

        wb.close()
        return extractions
